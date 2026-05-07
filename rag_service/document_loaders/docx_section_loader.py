import io
import re
import uuid

import unicodedata
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx.document import Document
from docx.oxml import CT_R, CT_Hyperlink
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from langchain.docstore.document import Document as LCDocument
from langchain.text_splitter import RecursiveCharacterTextSplitter, TextSplitter

from rag_service.constants import IMAGE_DOWNLOAD_URL_PROD, IMAGE_DOWNLOAD_URL_GAMMA
from rag_service.document_loaders.docx_loader import DocxLoader
from rag_service.env import ENV, EnvEnum
from rag_service.exceptions import ObsException
from rag_service.logger import Module, get_logger
from rag_service.models.generic.models import HeadAndContent
from rag_service.utils.his_util.obs_util import upload_file_as_bytes

logger = get_logger(module=Module.VECTORIZATION)


def iter_block_items(parent: Document):
    """获取Document对象的元素"""
    parent_elm = get_parent_elm(parent)

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def get_parent_elm(parent: Document):
    """获取元素内容"""
    if isinstance(parent, Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("对象类型错误, 应为Document类型或者_Cell类型")
    return parent_elm


def _handle_image(block, all_content, doc):
    if (('imagedata' in block._p.xml or 'pic:pic' in block._p.xml)
            and (match := re.search(r"rId\d+", block._p.xml))):
        r_id = match.group()
        image_part = doc.part.related_parts[r_id]
        ext = image_part.partname.split('.')[-1]
        blob = image_part.blob
        try:
            download_key = upload_file_as_bytes(str(uuid.uuid4()), blob)
            image_url = IMAGE_DOWNLOAD_URL_PROD if ENV == EnvEnum.PROD else IMAGE_DOWNLOAD_URL_GAMMA
            image_url = image_url.format(download_key, "image", ext)
            all_content[-1]["content"] += f"![]({image_url})"
        except ObsException as e:
            logger.error("Failed to upload content", e)


class DocxLoaderByHead(DocxLoader):
    def _parse_content(self) -> List[Dict[str, str]]:
        all_content = [{"title": "", "content": ""}]
        stack = []
        self.doc = docx.Document(self.doc_path)
        for block in iter_block_items(self.doc):
            if isinstance(block, Table):
                res = self._handle_table(block)
                all_content[-1]["content"] += res
            if not isinstance(block, Paragraph):
                continue
            handle_head = self.handle_paragraph_heading(all_content, block, stack)

            if block.style.name.startswith("Title"):
                all_content[-1]["title"] = block.text
                stack.append((0, block.text.strip()))
            elif not handle_head:
                if block.hyperlinks:
                    all_content[-1]["content"] += self.extract_hyperlink(block)
                else:
                    all_content[-1]["content"] += block.text

            _handle_image(block, all_content, self.doc)

            all_content[-1]["content"] += "\n"
        return all_content

    def load(self) -> List[LCDocument]:
        """将最小级别heading下的内容拼接生成Doc对象"""
        all_content = self._parse_content()
        docs = []
        for content in all_content:
            # 转化无意义特殊字符为标准字符
            plain_text = self._normalize(content["content"])
            # 过滤掉纯标题的document
            if len(plain_text) > 1:
                docs.append(
                    LCDocument(
                        page_content=f"{self._normalize(content['title'])} {plain_text}",
                        metadata={"source": self.doc_path},
                    )
                )
        return docs

    def load_and_split(self, text_splitter: Optional[TextSplitter] = None) -> List[LCDocument]:
        """将最小级别heading下的内容拼接生成Doc对象"""
        if text_splitter is None:
            _text_splitter: TextSplitter = RecursiveCharacterTextSplitter()
        else:
            _text_splitter = text_splitter

        all_content = self._parse_content()
        docs = []
        for content in all_content:
            # 转化无意义特殊字符为标准字符
            plain_text = self._normalize(content["content"])
            # 过滤掉纯标题的document
            if len(plain_text) > 1:
                # 按定长切分进行分组
                grouped_text = _text_splitter.split_text(plain_text)
                docs += [
                    LCDocument(
                        page_content=f"{self._normalize(content['title'])} {text}",
                        metadata={
                            "source": self.doc_path,
                            "headers": self._normalize(content['title']),
                            "header_contents": [content['content']]
                        }
                    )
                    for text in grouped_text
                ]
        return docs

    def get_head_contents(self):
        all_content = self._parse_content()
        head_contents: List[HeadAndContent] = []
        for content in all_content:
            # 转化无意义特殊字符为标准字符
            plain_text = self._normalize(content["content"])
            # 过滤掉纯标题的document
            if len(plain_text) > 1:
                head_contents.append(HeadAndContent(title=self._normalize(content["title"]), content=plain_text))
        return head_contents

    @classmethod
    def _normalize(cls, string: str) -> str:
        return unicodedata.normalize("NFKD", string).strip()

    @classmethod
    def handle_paragraph_heading(cls, all_content: List[Dict], block: Paragraph, stack: List[Tuple[Any, Any]]) -> bool:
        """处理Heading级别元素，并将上级标题拼接到本级标题中"""
        if block.style.name.startswith("Heading"):
            try:
                title_level = int(block.style.name.split()[-1])
            except Exception:
                return False
            while stack and stack[-1][0] >= title_level:
                stack.pop()
            if stack:
                parent_title = "".join(stack[-1][1])
                block.text = parent_title + "-" + block.text

            stack.append((title_level, block.text.strip()))
            all_content.append({"title": block.text, "content": " "})
            return True
        return False

    def extract_hyperlink(self, block):
        hyperlink_text = io.StringIO()
        try:
            for element in block.paragraph_format.element.inner_content_elements:
                hyperlink_text.write(self._handle_hyperlink_element(block, element))
            return hyperlink_text.getvalue()
        except Exception:
            return block.text or ""

    @classmethod
    def _handle_hyperlink_element(cls, block, element):
        if isinstance(element, CT_R):
            return element.text
        if isinstance(element, CT_Hyperlink):
            if element.text == block.part.rels[element.rId].target_ref:
                return element.text
            return f"{element.text} {block.part.rels[element.rId].target_ref}"
        return ""