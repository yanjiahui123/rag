from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from rag_service.document_loaders.parsed_blocks import BLOCK_TYPE_TABLE, SPLIT_POLICY_NO_SPLIT, ParsedBlock, ParsedDocument
from rag_service.document_loaders.structured_artifacts import STRUCTURED_EXCEL_ARTIFACTS_KEY
from rag_service.document_loaders.structured_loader import StructuredDocumentLoader
from rag_service.document_loaders.table.excel_parser import CellRange, ExcelTable, ExcelTableParser


class StructuredExcelLoader(StructuredDocumentLoader):
    def __init__(self, file_path: str, source: Optional[str] = None, chunk_size: int = 1500):
        super().__init__(file_path, source)
        self.chunk_size = chunk_size
        self.parser = ExcelTableParser()

    def parse_blocks(self) -> List[ParsedBlock]:
        return self.parse_to_document().to_blocks()

    def parse_to_document(self) -> ParsedDocument:
        sheets, merge_ranges = self._open_workbook_data()
        return self.parse_sheets(sheets, merge_ranges)

    def parse_sheets(
        self,
        sheets: Dict[str, List[List[Any]]],
        merge_ranges: Optional[Dict[str, List[CellRange]]] = None,
    ) -> ParsedDocument:
        blocks = []
        tables = []
        table_index = 0
        for sheet_name, grid in sheets.items():
            sheet_merges = (merge_ranges or {}).get(sheet_name, [])
            for table in self.parser.parse_grid(grid, sheet_name, Path(self.source).name, sheet_merges):
                table.table_id = self._table_id(table_index)
                table_index += 1
                tables.append(table)
                blocks.extend(self._table_blocks(table))
        return ParsedDocument(blocks=blocks, metadata=self._metadata(blocks, tables))

    def _open_sheets(self) -> Dict[str, List[List[Any]]]:
        return self._open_workbook_data()[0]

    def _open_workbook_data(self) -> Tuple[Dict[str, List[List[Any]]], Dict[str, List[CellRange]]]:
        from openpyxl.reader.excel import load_workbook

        workbook = load_workbook(self.file_path, data_only=True)
        sheets = {
            worksheet.title: [list(row) for row in worksheet.iter_rows(values_only=True)]
            for worksheet in workbook.worksheets
        }
        merge_ranges = {worksheet.title: self._merge_ranges(worksheet) for worksheet in workbook.worksheets}
        return sheets, merge_ranges

    @staticmethod
    def _merge_ranges(worksheet) -> List[CellRange]:
        return [
            (cell_range.min_row - 1, cell_range.min_col - 1, cell_range.max_row, cell_range.max_col)
            for cell_range in getattr(worksheet.merged_cells, "ranges", [])
        ]

    def _table_blocks(self, table: ExcelTable) -> List[ParsedBlock]:
        blocks = []
        for chunk_index, chunk in enumerate(table.to_llm_chunks(max_chars=self.chunk_size)):
            blocks.append(
                ParsedBlock(
                    text=chunk.text,
                    metadata=self._chunk_metadata(table, chunk, chunk_index),
                    block_type=BLOCK_TYPE_TABLE,
                    split_policy=SPLIT_POLICY_NO_SPLIT,
                )
            )
        return blocks

    def _chunk_metadata(self, table: ExcelTable, chunk, chunk_index: int) -> Dict[str, Any]:
        return {
            "source": self.source,
            "loader": "structured_excel",
            "block_id": f"{table.table_id}_chunk_{chunk_index + 1:03d}",
            "title": table.title,
            "table_id": table.table_id,
            "sheet_name": table.sheet_name,
            "cell_range": table.cell_range,
            "row_range": chunk.row_range,
            "oversized_row": chunk.oversized_row,
            "table": self._table_summary(table),
        }

    @staticmethod
    def _table_summary(table: ExcelTable) -> Dict[str, Any]:
        return {
            "table_id": table.table_id,
            "title": table.title,
            "source_type": "excel_table",
            "sheet_name": table.sheet_name,
            "cell_range": table.cell_range,
            "description": table.description,
            "flatten_headers": list(table.flatten_headers),
            "row_count": len(table.rows),
            "col_count": len(table.flatten_headers),
        }

    def _metadata(self, blocks: List[ParsedBlock], tables: List[ExcelTable]) -> Dict[str, Any]:
        return {
            "source": self.source,
            STRUCTURED_EXCEL_ARTIFACTS_KEY: {
                "document_markdown": self._document_markdown(tables),
                "tables": [self._table_artifact(table) for table in tables],
            },
        }

    def _document_markdown(self, tables: List[ExcelTable]) -> str:
        return "\n\n".join(self._table_llm_text(table) for table in tables)

    def _table_artifact(self, table: ExcelTable) -> Dict[str, Any]:
        return {
            "table_id": table.table_id,
            "html": table.display_html,
            "json": table.to_artifact_dict(),
            "llm_markdown": self._table_llm_text(table),
        }

    @staticmethod
    def _table_llm_text(table: ExcelTable) -> str:
        chunks = table.to_llm_chunks(max_chars=10**9)
        return "\n\n".join(chunk.text for chunk in chunks)

    @staticmethod
    def _table_id(table_index: int) -> str:
        return f"table_{table_index + 1:03d}"
