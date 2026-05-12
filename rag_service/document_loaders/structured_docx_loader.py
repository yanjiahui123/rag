from __future__ import annotations

import re
from typing import Any, Dict, Iterable, List, Optional

from pydantic import BaseModel, Field

from rag_service.document_loaders.image_markdown import (
    ImageMarkdown,
    image_extension_from_partname,
    upload_image_bytes,
)
from rag_service.document_loaders.parsed_blocks import (
    BLOCK_TYPE_TABLE,
    BLOCK_TYPE_TEXT,
    SPLIT_POLICY_MARKDOWN_HEADINGS,
    SPLIT_POLICY_NO_SPLIT,
    ParsedBlock,
    ParsedDocument,
)
from rag_service.document_loaders.structured_artifacts import STRUCTURED_DOCX_ARTIFACTS_KEY
from rag_service.document_loaders.structured_loader import StructuredDocumentLoader
from rag_service.document_loaders.table.models import TableBlock
from rag_service.document_loaders.table.docx_parser import DocxTableParser

IMAGE_RELATION_PATTERN = re.compile(r'\br:(?:embed|id)="([^"]+)"')
IMAGE_MARKERS = ("pic:pic", "imagedata", "a:blip")


class DocxMarkdownElement(BaseModel):
    text: str
    block_index: int
    headers: List[str] = Field(default_factory=list)
    heading_level: Optional[int] = None
    table_id: Optional[str] = None
    table_block: Optional[TableBlock] = None


class StructuredDocxLoader(StructuredDocumentLoader):
    def __init__(self, file_path: str, source: Optional[str] = None, table_parser: Optional[DocxTableParser] = None):
        super().__init__(file_path, source)
        self.table_parser = table_parser or DocxTableParser()
        self.image_object_keys: List[str] = []

    def parse_blocks(self) -> List[ParsedBlock]:
        return self.parse_to_document().to_blocks()

    def parse_to_document(self) -> ParsedDocument:
        return self.parse_document(self._open_document())

    def parse_document(self, document: Any) -> ParsedDocument:
        self.image_object_keys = []
        elements = self._document_to_markdown_elements(document)
        if not elements:
            return ParsedDocument(metadata=self._base_metadata([], 0))
        return self._document_from_elements(elements)

    def _document_to_markdown_elements(self, document: Any) -> List[DocxMarkdownElement]:
        elements = []
        headers = []
        table_index = 0
        for block_index, block in enumerate(self._iter_blocks(document)):
            if self._is_table(block):
                elements.append(self._table_element(block, headers, block_index, table_index))
                table_index += 1
                continue
            headers = self._append_paragraph_element(elements, block, headers, block_index, document)
        return elements

    def _append_paragraph_element(self, elements, block, headers, block_index, document=None):
        text = self._paragraph_text(block)
        image_markdown = self._paragraph_image_markdown(block, document)
        if not text and not image_markdown:
            return headers
        heading_level = self._heading_level(block) if text else None
        if heading_level is None:
            elements.append(self._text_element(_combine_text_and_images(text, image_markdown), headers, block_index))
            return headers
        next_headers = self._replace_header(headers, heading_level, text)
        elements.append(self._heading_element(text, heading_level, next_headers, block_index))
        if image_markdown:
            elements.append(self._text_element("\n".join(image_markdown), next_headers, block_index))
        return next_headers

    def _open_document(self) -> Any:
        import docx

        return docx.Document(self.file_path)

    def _iter_blocks(self, document: Any) -> Iterable[Any]:
        if hasattr(document, "blocks"):
            return document.blocks
        from rag_service.document_loaders.docx_section_loader import iter_block_items

        return iter_block_items(document)

    @staticmethod
    def _is_table(block: Any) -> bool:
        return hasattr(block, "rows")

    @staticmethod
    def _paragraph_text(block: Any) -> str:
        return str(getattr(block, "text", "") or "").strip()

    def _paragraph_image_markdown(self, block: Any, document: Any) -> List[str]:
        xml = str(getattr(getattr(block, "_p", None), "xml", "") or "")
        if not _contains_image_marker(xml):
            return []
        related_parts = _related_parts(document, block)
        image_links = []
        for relation_id in _image_relation_ids(xml):
            image_part = related_parts.get(relation_id)
            if not _is_image_part(image_part):
                continue
            image = _upload_image_part(image_part)
            if image.markdown:
                image_links.append(image.markdown)
            if image.object_key:
                self.image_object_keys.append(image.object_key)
        return image_links

    def _heading_level(self, block: Any) -> Optional[int]:
        style_name = getattr(getattr(block, "style", None), "name", "")
        if style_name.startswith("Title"):
            return 1
        if not style_name.startswith("Heading"):
            return None
        return self._parse_heading_level(style_name)

    @staticmethod
    def _parse_heading_level(style_name: str) -> Optional[int]:
        try:
            return int(style_name.split()[-1])
        except Exception:
            return None

    @staticmethod
    def _replace_header(headers: List[str], level: int, text: str) -> List[str]:
        index = max(level - 1, 0)
        next_headers = list(headers[:index])
        while len(next_headers) < index:
            next_headers.append("")
        next_headers.append(text)
        return next_headers

    @staticmethod
    def _heading_element(text: str, level: int, headers: List[str], block_index: int) -> DocxMarkdownElement:
        marker = "#" * max(level, 1)
        return DocxMarkdownElement(
            text=f"{marker} {text}",
            block_index=block_index,
            headers=list(headers),
            heading_level=level,
        )

    @staticmethod
    def _text_element(text: str, headers: List[str], block_index: int) -> DocxMarkdownElement:
        return DocxMarkdownElement(text=text, block_index=block_index, headers=list(headers))

    def _table_element(self, table: Any, headers: List[str], block_index: int, table_index: int) -> DocxMarkdownElement:
        metadata = self._base_metadata(headers, block_index)
        metadata["table_index"] = table_index
        metadata["table_id"] = self._table_id(table_index)
        metadata["title"] = headers[-1] if headers else ""
        table_block = self.table_parser.parse(table, metadata=metadata)
        return DocxMarkdownElement(
            text=self._table_search_text(table_block, headers),
            block_index=block_index,
            headers=list(headers),
            table_id=metadata["table_id"],
            table_block=table_block,
        )

    def _base_metadata(self, headers: List[str], block_index: int) -> dict:
        return {
            "source": self.source,
            "headers": list(headers),
            "block_index": block_index,
            "loader": "structured_docx",
        }

    def _document_from_elements(self, elements: List[DocxMarkdownElement]) -> ParsedDocument:
        blocks = self._blocks_from_elements(elements)
        return ParsedDocument(blocks=blocks, metadata=self._document_metadata(elements, blocks))

    def _document_metadata(self, elements: List[DocxMarkdownElement], blocks: List[ParsedBlock]) -> Dict[str, Any]:
        metadata = self._base_metadata([], 0)
        metadata[STRUCTURED_DOCX_ARTIFACTS_KEY] = {
            "document_markdown": self._document_markdown(elements),
            "manifest": self._manifest(blocks),
            "tables": self._table_artifacts(elements),
            "image_object_keys": list(dict.fromkeys(self.image_object_keys)),
        }
        return metadata

    def _blocks_from_elements(self, elements: List[DocxMarkdownElement]) -> List[ParsedBlock]:
        blocks = []
        text_elements = []
        for element in elements:
            if element.table_block:
                self._append_text_block(blocks, text_elements)
                text_elements = []
                blocks.append(self._table_parsed_block(element))
            else:
                if element.heading_level is not None and text_elements:
                    self._append_text_block(blocks, text_elements)
                    text_elements = []
                text_elements.append(element)
        self._append_text_block(blocks, text_elements)
        return blocks

    def _append_text_block(self, blocks: List[ParsedBlock], text_elements: List[DocxMarkdownElement]) -> None:
        if not text_elements:
            return
        text = "\n\n".join(element.text for element in text_elements if element.text).strip()
        if text:
            blocks.append(
                ParsedBlock(
                    text=text,
                    metadata=self._text_block_metadata(text_elements),
                    split_policy=SPLIT_POLICY_MARKDOWN_HEADINGS,
                )
            )

    def _text_block_metadata(self, text_elements: List[DocxMarkdownElement]) -> Dict[str, Any]:
        first_element = text_elements[0]
        headers = list(text_elements[-1].headers)
        metadata = self._base_metadata(headers, first_element.block_index)
        metadata["block_id"] = self._block_id(first_element.block_index)
        metadata["title"] = headers[-1] if headers else ""
        return metadata

    def _table_parsed_block(self, element: DocxMarkdownElement) -> ParsedBlock:
        return ParsedBlock(
            text=element.text,
            metadata=self._table_block_metadata(element),
            block_type=BLOCK_TYPE_TABLE,
            split_policy=SPLIT_POLICY_NO_SPLIT,
        )

    def _table_block_metadata(self, element: DocxMarkdownElement) -> Dict[str, Any]:
        table_block = element.table_block
        metadata = self._base_metadata(element.headers, element.block_index)
        metadata["block_id"] = self._block_id(element.block_index)
        metadata["title"] = table_block.title
        metadata["table_id"] = element.table_id
        metadata["table"] = self._table_summary(table_block, element.table_id)
        return metadata

    @staticmethod
    def _table_summary(table_block: TableBlock, table_id: str) -> Dict[str, Any]:
        return {
            "table_id": table_id,
            "title": table_block.title,
            "source_type": table_block.source_type,
            "flatten_headers": list(table_block.headers),
            "row_count": len(table_block.rows),
            "col_count": len(table_block.headers),
        }

    def _table_search_text(self, table_block: TableBlock, headers: List[str]) -> str:
        return table_block.to_llm_text(section_headers=headers)

    @staticmethod
    def _document_markdown(elements: List[DocxMarkdownElement]) -> str:
        return "\n\n".join(element.text for element in elements if element.text).strip()

    @staticmethod
    def _manifest(blocks: List[ParsedBlock]) -> Dict[str, Any]:
        return {
            "blocks": [
                {
                    "block_id": block.metadata.get("block_id"),
                    "type": block.metadata.get("block_type", BLOCK_TYPE_TEXT),
                    "title": block.metadata.get("title", ""),
                    "headers": block.metadata.get("headers", []),
                    "table_id": block.metadata.get("table_id"),
                }
                for block in blocks
            ]
        }

    @staticmethod
    def _table_artifacts(elements: List[DocxMarkdownElement]) -> List[Dict[str, Any]]:
        return [
            {
                "table_id": element.table_id,
                "html": element.table_block.display_html or "",
                "json": element.table_block.to_artifact_dict(),
                "llm_markdown": element.table_block.to_llm_text(section_headers=element.headers),
            }
            for element in elements
            if element.table_block
        ]

    @staticmethod
    def _block_id(block_index: int) -> str:
        return f"block_{block_index + 1:03d}"

    @staticmethod
    def _table_id(table_index: int) -> str:
        return f"table_{table_index + 1:03d}"


def _contains_image_marker(xml: str) -> bool:
    return any(marker in xml for marker in IMAGE_MARKERS)


def _image_relation_ids(xml: str) -> List[str]:
    relation_ids = []
    seen = set()
    for relation_id in IMAGE_RELATION_PATTERN.findall(xml):
        if relation_id in seen:
            continue
        seen.add(relation_id)
        relation_ids.append(relation_id)
    return relation_ids


def _related_parts(document: Any, block: Any) -> Dict[str, Any]:
    for owner in (document, block):
        part = getattr(owner, "part", None)
        related_parts = getattr(part, "related_parts", None)
        if isinstance(related_parts, dict):
            return related_parts
    return {}


def _is_image_part(image_part: Any) -> bool:
    if image_part is None:
        return False
    content_type = str(getattr(image_part, "content_type", "") or "")
    if content_type.startswith("image/"):
        return True
    partname = str(getattr(image_part, "partname", "") or "").lower()
    return partname.endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".svg"))


def _upload_image_part(image_part: Any) -> ImageMarkdown:
    extension = image_extension_from_partname(getattr(image_part, "partname", ""))
    return upload_image_bytes(getattr(image_part, "blob", b""), extension)


def _combine_text_and_images(text: str, image_markdown: List[str]) -> str:
    parts = []
    if text:
        parts.append(text)
    parts.extend(image_markdown)
    return "\n".join(parts).strip()
