import itertools
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any, List, Tuple, Union

import pandas as pd
from langchain.docstore.document import Document
from langchain_community.document_loaders.base import BaseLoader
from openpyxl.reader.excel import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pandas import DataFrame, ExcelFile, Series
from xlrd import open_workbook
from xlrd.sheet import Sheet

from rag_service.constants import LOADER_EXCEL_SIZE_LIMIT
from rag_service.document_loaders.loader_utils import build_markdown_table_string
from rag_service.logger import Module, get_logger
from rag_service.rag_app.api_exceptions import FileTooLarge
from rag_service.text_splitters.markdown_table_splitter import MarkdownTableSplitter

logger = get_logger(module=Module.VECTORIZATION)
_CHUNK_SIZE = 100000


def is_file_over_input_size(file_path: str, max_size: int):
    path = Path(file_path)
    if not path.exists():
        return False
    file_size = path.stat().st_size
    return file_size > max_size


class _BaseExcelLoader(BaseLoader, ABC):
    def __init__(self, file_path: str):
        self.excel_path = file_path

    def load(self) -> List[Document]:
        """Load documents."""
        if is_file_over_input_size(self.excel_path, LOADER_EXCEL_SIZE_LIMIT):
            logger.error("文件%s超过规定大小", self.excel_path)
            raise FileTooLarge(f"文件{self.excel_path}超过规定大小")
        all_text = self._handle_excel()
        return [Document(page_content=one_text, metadata={"source": self.excel_path}) for one_text in all_text]

    def _handle_excel(self) -> List[str]:
        if not Path(self.excel_path).exists():
            logger.error("文件%s不存在", self.excel_path)

        excel = self._get_excel()
        return itertools.chain.from_iterable([
            self._handle_sheet(excel, sheet_name) for sheet_name in excel.sheet_names
        ])

    @abstractmethod
    def _get_excel(self) -> ExcelFile: ...

    def _handle_sheet(self, excel: ExcelFile, sheet_name: str) -> List[str]:
        sheet: Union[Worksheet, Sheet] = excel.book[sheet_name]
        merged_cells = self._handle_merged_cells(sheet)

        chunk_count = 0
        # row number in 1-based index, and we assume the 1st row is a header, so the content row number starts from 2.
        chunk_start_row = 2
        chunk_end_row = chunk_start_row + _CHUNK_SIZE - 1

        result = []
        while True:
            data_frame_chunk: DataFrame = excel.parse(
                sheet_name,
                keep_default_na=False,
                nrows=_CHUNK_SIZE,
                skiprows=lambda _: _ in range(1, chunk_start_row - 1),
            )
            if data_frame_chunk.empty:
                break

            for top_col, top_row, bottom_col, bottom_row, base_value in merged_cells:
                if bottom_row < chunk_start_row or top_row > chunk_end_row:
                    continue
                if base_value is not None:
                    df_start_row = max(top_row, chunk_start_row) - 2 - _CHUNK_SIZE * chunk_count
                    df_end_row = min(bottom_row, chunk_end_row) - 1 - _CHUNK_SIZE * chunk_count
                    data_frame_chunk.iloc[df_start_row:df_end_row, top_col:bottom_col] = base_value
            markdown_table_list = self._excel_to_markdown_table(data_frame_chunk, self.excel_path, sheet_name)
            result.extend(markdown_table_list)
            chunk_start_row += _CHUNK_SIZE
            chunk_end_row += _CHUNK_SIZE
            chunk_count += 1
        return result

    @abstractmethod
    def _handle_merged_cells(self, sheet: Union[Worksheet, Sheet]) -> List[Tuple[int, int, int, int, Any]]: ...

    @abstractmethod
    def _handle_row(self, sheet: Union[Worksheet, Sheet], row: Series) -> str: ...

    @abstractmethod
    def _excel_to_markdown_table(self, data_frame_chunk: DataFrame, excel_path: str, sheet_name: str) -> List[str]:
        ...


class _BaseXlsxLoader(_BaseExcelLoader, ABC):
    def _get_excel(self) -> ExcelFile:
        return pd.ExcelFile(load_workbook(self.excel_path), engine="openpyxl")

    def _handle_merged_cells(self, sheet: Union[Worksheet, Sheet]) -> List[Tuple[int, int, int, int, Any]]:
        merged_cells = []
        for item in sheet.merged_cells:
            top_col, top_row, bottom_col, bottom_row = item.bounds
            base_value = item.start_cell.value
            merged_cells.append((top_col - 1, top_row, bottom_col, bottom_row, base_value))
        return merged_cells


class _BaseXlsLoader(_BaseExcelLoader, ABC):
    def _get_excel(self) -> ExcelFile:
        return pd.ExcelFile(open_workbook(self.excel_path, formatting_info=True), engine="xlrd")

    def _handle_merged_cells(self, sheet: Union[Worksheet, Sheet]) -> List[Tuple[int, int, int, int, Any]]:
        merged_cells = []
        for top_row, bottom_row, top_col, bottom_col in sheet.merged_cells:
            base_value = sheet.cell_value(top_row, top_col)
            merged_cells.append((top_col, top_row + 1, bottom_col, bottom_row, base_value))
        return merged_cells


class _RowConcatenatorMixin:
    @classmethod
    def concat_row(cls, filename: str, sheet_name: str, row: Series) -> str:
        row_text = f"file_name: {filename}, sheet_name: {sheet_name}, "
        for column_name, value in row.items():
            if not str(column_name).startswith("Unnamed"):
                row_text += f"{column_name}: {value}, "
        return row_text

    @classmethod
    def header_and_row_to_markdown(cls, data_frame_chunk: DataFrame, excel_path: str, sheet_name: str) -> List[str]:
        row_text = f"file_name: {Path(excel_path).name}；sheet_name: {sheet_name}\n"
        result = []
        for _, row in data_frame_chunk.iterrows():
            headers = []
            row_value = []
            for column_name, value in row.items():
                if not str(column_name).startswith("Unnamed"):
                    headers.append(str(column_name).replace('\n', ' ').replace('|', '\|').strip())
                    row_value.append(str(value).replace('\n', ' ').replace('|', '\|').strip())
            markdown_header_and_row = build_markdown_table_string(headers, [row_value])
            result.append(row_text + markdown_header_and_row)
        return result


class XlsxLoader(_BaseXlsxLoader, _RowConcatenatorMixin):
    def _excel_to_markdown_table(self, data_frame_chunk: DataFrame, excel_path: str, sheet_name: str) -> List[str]:
        return self.header_and_row_to_markdown(data_frame_chunk, excel_path, sheet_name)

    def _handle_row(self, sheet: Worksheet, row: Series) -> str:
        return self.concat_row(Path(self.excel_path).name, sheet.title, row)


class XlsLoader(_BaseXlsLoader, _RowConcatenatorMixin):
    def _excel_to_markdown_table(self, data_frame_chunk: DataFrame, excel_path: str, sheet_name: str) -> List[str]:
        return self.header_and_row_to_markdown(data_frame_chunk, excel_path, sheet_name)

    def _handle_row(self, sheet: Sheet, row: Series) -> str:
        return self.concat_row(Path(self.excel_path).name, sheet.name, row)


class _MarkdownConverterMixin:
    @classmethod
    def convert_row_to_markdown_table(cls, filename: str, sheet_name: str, row: Series) -> str:
        data = {"file_name": Path(filename).name, "sheet_name": sheet_name}
        for column_name, value in row.items():
            if not str(column_name).startswith("Unnamed"):
                data[str(column_name)] = value if not isinstance(value, str) else value.replace("\n", "<br>")
        df = pd.DataFrame.from_dict(data=data, orient="index", columns=[""])
        return df.to_markdown(tablefmt="github", numalign=None, stralign=None)

    @classmethod
    def entire_table_to_markdown(cls, data_frame_chunk: DataFrame, excel_path: str, sheet_name: str) -> List[str]:
        headers = [str(col).replace('\n', ' ').replace('|', '\|').strip()
                   for col in data_frame_chunk.columns.tolist()]
        rows = [[str(item).replace('\n', ' ').replace('|', '\|').strip() for item in row]
                for row in data_frame_chunk.values.tolist()]
        markdown_table = build_markdown_table_string(headers, rows)
        splitter = MarkdownTableSplitter()
        split_markdown_tables = splitter.split_document_by_tables(markdown_table, excel_path)
        file_name = Path(excel_path).name
        return [f"file_name：{file_name}；sheet_name：{sheet_name}\n{split_markdown_table}"
                for split_markdown_table in split_markdown_tables]


class XlsxToMarkdownLoader(_BaseXlsxLoader, _MarkdownConverterMixin):
    def _excel_to_markdown_table(self, data_frame_chunk: DataFrame, excel_path: str, sheet_name: str) -> List[str]:
        return self.entire_table_to_markdown(data_frame_chunk, excel_path, sheet_name)

    def _handle_row(self, sheet: Worksheet, row: Series) -> str:
        return self.convert_row_to_markdown_table(Path(self.excel_path).name, sheet.title, row)


class XlsToMarkdownLoader(_BaseXlsLoader, _MarkdownConverterMixin):
    def _excel_to_markdown_table(self, data_frame_chunk: DataFrame, excel_path: str, sheet_name: str) -> List[str]:
        return self.entire_table_to_markdown(data_frame_chunk, excel_path, sheet_name)

    def _handle_row(self, sheet: Sheet, row: Series) -> str:
        return self.convert_row_to_markdown_table(Path(self.excel_path).name, sheet.name, row)